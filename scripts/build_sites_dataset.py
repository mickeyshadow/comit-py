"""Generate datasets/sites.csv - THE V2 UNIVERSE: all UK industry, in
transparent layers, anchored to the national territorial total.

Layers (the `layer` column makes reporting by coverage class trivial):
  ets_named        top-100 ETS industrial installations by 2025 verified
                   emissions (UK ETS registry) - traded
  ets_tail         per-sector aggregates of the remaining ETS installations
                   - traded (v1 mislabelled these non-traded; fixed)
  nontraded_named  NAEI non-traded point sources >=10 kt (COMIT template's
                   public NAEI table, 2023 vintage scaled to 2025 by
                   crosswalk sector ratios) - untraded carbon price
  nontraded_tail   per-sector aggregate of smaller non-traded point sources
  diffuse          the remainder to the national anchor: DESNZ 2024
                   provisional territorial industry = 13% of 371 MtCO2e
                   ~ 48 Mt; scaled ~ -3% to 2025 -> 46.5 Mt anchor

Double-count guard: NAEI labels parts of ETS complexes non-traded
(Scunthorpe Sinter, Port Talbot Power Station). Non-traded rows sharing a
distinctive name token with a named ETS site are excluded and logged.

The hindcast universe (datasets_backcast/) deliberately stays the
ETS-verified core - skill remains defined where outturn exists.
"""
import csv
import re
from collections import defaultdict

import openpyxl

ETS = r"D:\comit-harness\data\uk_ets_compliance_2021_2025.xlsx"
CW = r"D:\comit-harness\data\ets_naei_crosswalk.csv"
NAEI_XLSX = (r"D:\comit-harness\comit\data_template_archive"
             r"\comit_input_1_4_0_public_updated.xlsx")
OUT = r"D:\comit-py\datasets\sites.csv"

SRC_ETS = "UK ETS compliance report 2021-2025 (registry, pub 22 Jun 2026)"
SRC_NAEI = "NAEI point sources via COMIT template (2023 vintage, OGL)"
RET = "2026-08-16"

ANCHOR_2025_KT = 46_500  # DESNZ 2024 territorial industry ~48 Mt (13% of 371), ~-3% to 2025
INTENSITY = {"steel": 1900.0, "cement": 750.0}
GENERIC = set("""cement steel lime glass paper mill mills refinery blast
furnaces furnace kiln works plant station power chemical chemicals limited
ltd uk group energy factory sugar""".split())


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def tokens(name):
    s = re.sub(r"[^a-z0-9 ]", " ", str(name).lower())
    return {t for t in s.split() if len(t) >= 5 and t not in GENERIC}


def sector_of_nace(nace, desc):
    d = (desc or "").lower()
    if nace.startswith("241") or "basic iron and steel" in d:
        return "steel"
    if nace.startswith("2351") or "cement" in d:
        return "cement"
    return "heat"


def sector_of_ipm(ipm):
    s = str(ipm)
    if "ron & steel" in s:
        return "steel"
    if s == "Cement":
        return "cement"
    return "heat"


def demand_and_tech(sector, kt):
    if sector == "steel":
        return kt / INTENSITY["steel"], "bof"
    if sector == "cement":
        return kt / INTENSITY["cement"], "cement_kiln"
    return kt / 51.2 * 0.9, "gas_boiler"


# ---- crosswalk: permit -> plant (for clusters) and sector vintage ratios ----
cw_permit = {}
cw_rows = []
with open(CW, encoding="utf-8") as fh:
    for row in csv.DictReader(fh):
        cw_rows.append(row)
        for part in row["permits"].split(";"):
            if "x" in part:
                permit, frac = part.rsplit("x", 1)
                if permit not in cw_permit or float(frac) > cw_permit[permit][1]:
                    cw_permit[permit] = (row["plant_id"], float(frac))

ratio_num = defaultdict(float)
ratio_den = defaultdict(float)
for row in cw_rows:
    sec = sector_of_ipm(row["sector"])
    ratio_num[sec] += float(row["ets_2025_t"])
    ratio_den[sec] += float(row["ets_2023_t"])
vintage_ratio = {s: (ratio_num[s] / ratio_den[s] if ratio_den[s] else 1.0)
                 for s in ("heat", "steel", "cement")}

# ---- NAEI table: clusters + non-traded point sources ------------------------
_naei = openpyxl.load_workbook(NAEI_XLSX, read_only=True)
_nr = list(_naei["NAEI_df_clean_2023_revised"].iter_rows(values_only=True))
_hi = next(i for i, r in enumerate(_nr) if r and "PlantID" in [str(c) for c in r])
_nh = {h: i for i, h in enumerate(_nr[_hi])}
plant_cluster = {str(r[_nh["PlantID"]]): str(r[_nh["cluster_location"]])
                 for r in _nr[_hi + 1:] if r and r[_nh["PlantID"]]}

nontraded = []
for r in _nr[_hi + 1:]:
    if not r or not r[_nh["PlantID"]] or r[_nh["traded_flag"]] != "Non-traded":
        continue
    kt23 = (num(r[_nh["Emissions_tco2e"]]) or 0) / 1e3
    if kt23 <= 0:
        continue
    sec = sector_of_ipm(r[_nh["IPM_sector"]])
    c = str(r[_nh["cluster_location"]])
    nontraded.append({
        "name": str(r[_nh["Site"]]),
        "sector": sec,
        "kt25": kt23 * vintage_ratio[sec],
        "cluster": "Dispersed" if c == "Not in cluster" else c,
    })


def cluster_of(permit):
    hit = cw_permit.get(str(permit))
    if not hit:
        return "Unmapped"
    c = plant_cluster.get(hit[0], "Unmapped")
    return "Dispersed" if c == "Not in cluster" else c


# ---- ETS layer --------------------------------------------------------------
wb = openpyxl.load_workbook(ETS, read_only=True)
rows = list(wb["Data"].iter_rows(values_only=True))
hdr = {h: i for i, h in enumerate(rows[0])}
ets = []
for r in rows[1:]:
    if r[hdr["Account type"]] != "OPERATOR_HOLDING_ACCOUNT":
        continue
    nace = str(r[hdr["NACE"]] or "")
    desc = str(r[hdr["NACE Description"]] or "")
    if nace.startswith("35") or "electricity" in desc.lower() \
            or "extraction" in desc.lower():
        continue
    e25 = num(r[hdr["Recorded emissions 2025"]])
    ly = num(r[hdr["Last Year of Operation"]])
    if e25 is None or e25 < 1000 or (ly and ly <= 2025):
        continue
    ets.append({
        "name": (r[hdr["Installation name"]] or r[hdr["Account Holder Name"]]),
        "sector": sector_of_nace(nace, desc),
        "kt25": e25 / 1e3,
        "cluster": cluster_of(r[hdr["Permit ID or Monitoring plan ID"]]),
    })
ets.sort(key=lambda e: -e["kt25"])
top, tail = ets[:100], ets[100:]

# ---- double-count guard -----------------------------------------------------
ets_tokens = set()
for e in top:
    ets_tokens |= tokens(e["name"])
kept, excluded = [], []
for s in nontraded:
    if tokens(s["name"]) & ets_tokens:
        excluded.append(s)
    else:
        kept.append(s)
print(f"double-count guard: excluded {len(excluded)} non-traded rows "
      f"({sum(s['kt25'] for s in excluded):.0f} kt) sharing names with ETS "
      f"sites, e.g. "
      f"{[s['name'] for s in sorted(excluded, key=lambda x: -x['kt25'])[:4]]}")

nt_named = [s for s in kept if s["kt25"] >= 10]
nt_tail_kt = defaultdict(float)
for s in kept:
    if s["kt25"] < 10:
        nt_tail_kt[s["sector"]] += s["kt25"]

# ---- inertia factors (band base x size gradient, per layer) -----------------
BAND_BASE = {"large": 0.8, "mid": 1.0, "small": 1.3}


def band_of(kt):
    return "large" if kt > 500 else ("mid" if kt > 50 else "small")


def with_factors(entries):
    by_band = defaultdict(list)
    for e in entries:
        by_band[band_of(e["kt25"])].append(e)
    fx = {}
    for band, group in by_band.items():
        group.sort(key=lambda x: -x["kt25"])
        n = len(group)
        for i, e in enumerate(group):
            grad = 0.7 + 0.6 * (i / (n - 1) if n > 1 else 0.5)
            fx[e["name"]] = round(BAND_BASE[band] * grad, 3)
    return fx


fx_ets = with_factors(top)
fx_nt = with_factors(nt_named)

# ---- assemble ---------------------------------------------------------------
out = []
_seen: dict[str, int] = {}


def unique_name(name):
    """NAEI non-traded names collide (several distinct 'Trafford Park'
    rows); duplicate site names would corrupt the LP variable index."""
    n = _seen.get(name, 0) + 1
    _seen[name] = n
    return name if n == 1 else f"{name} ({n})"


def add_row(name, sector, band, traded, kt, tech_default, factor, cluster,
            layer, src, basis):
    d, tech = demand_and_tech(sector, kt)
    commodity = sector if sector in ("steel", "cement") else "heat"
    out.append([unique_name(name), sector, band, traded, commodity,
                round(d, 4), tech, round(d / 0.9, 4), factor, cluster, layer,
                src, RET, basis])


for e in top:
    add_row(e["name"], e["sector"], band_of(e["kt25"]), True, e["kt25"],
            None, fx_ets[e["name"]], e["cluster"], "ets_named", SRC_ETS,
            "derived")

tail_kt = defaultdict(float)
for e in tail:
    tail_kt[e["sector"]] += e["kt25"]
for sector, kt in sorted(tail_kt.items()):
    add_row(f"ets_tail_{sector}", sector, "small", True, kt, None, 1.3,
            "Dispersed", "ets_tail",
            SRC_ETS + " (sub-100 aggregate; traded - v1 flag fixed)", "derived")

for s in nt_named:
    add_row(s["name"], s["sector"], band_of(s["kt25"]), False, s["kt25"],
            None, fx_nt[s["name"]], s["cluster"], "nontraded_named",
            SRC_NAEI + f" x{vintage_ratio[s['sector']]:.2f} to 2025", "derived")

for sector, kt in sorted(nt_tail_kt.items()):
    add_row(f"nontraded_tail_{sector}", sector, "small", False, kt, None, 1.3,
            "Dispersed", "nontraded_tail", SRC_NAEI + " (sub-10kt aggregate)",
            "derived")

point_total = (sum(e["kt25"] for e in top) + sum(tail_kt.values())
               + sum(s["kt25"] for s in nt_named) + sum(nt_tail_kt.values()))
diffuse_kt = max(0.0, ANCHOR_2025_KT - point_total)
add_row("diffuse_industry", "heat", "small", False, diffuse_kt, None, 1.3,
        "Dispersed", "diffuse",
        "remainder to DESNZ 2024 territorial industry anchor (13% of 371 Mt "
        "~48 Mt; -3% to 2025 = 46.5 Mt)", "derived")

with open(OUT, "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh)
    w.writerow(["site", "sector", "band", "traded", "commodity", "demand_pj",
                "incumbent_tech", "start_capacity", "inertia_factor",
                "cluster", "layer", "source", "retrieved", "basis"])
    w.writerows(out)

by_layer = defaultdict(float)
for r in out:
    by_layer[r[10]] += 0  # count via names below
print(f"wrote {OUT}: {len(out)} sites | layers: "
      f"ets_named {len(top)}, ets_tail {len(tail_kt)}, "
      f"nontraded_named {len(nt_named)}, nontraded_tail {len(nt_tail_kt)}, "
      f"diffuse 1")
print(f"universe 2025: point {point_total/1e3:.1f} Mt + diffuse "
      f"{diffuse_kt/1e3:.1f} Mt = {ANCHOR_2025_KT/1e3:.1f} Mt anchor")
